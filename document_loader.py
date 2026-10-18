"""
document_loader.py  (Phase 3)
PDF / DOCX / TXT / CSV documents se text nikalna, saaf karna, aur metadata jodna.

Har document "records" ki list ban jaata hai. Ek record = ek page ka text + metadata:
{
    "filename":   "EM-01_Emergency_Admission_Policy.pdf",
    "doc_id":     "EM-01",
    "title":      "Emergency Admission Policy",
    "department": "Emergency",
    "file_type":  "pdf",
    "page":       1,
    "text":       "...saaf kiya hua text..."
}
"""

import re
from pathlib import Path

import fitz                      # PyMuPDF: PDF padhne ke liye
import pandas as pd              # CSV padhne ke liye
from docx import Document        # python-docx: DOCX padhne ke liye

# Filename ke prefix se department pata karte hain
DEPARTMENTS = {
    "EM": "Emergency",
    "ADM": "Admissions",
    "PS": "Patient Safety",
    "PH": "Pharmacy",
    "NR": "Nursing",
    "BF": "Billing & Finance",
}

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}

# Disclaimer + PDF footer ("... | EM-01 | Page 1 of 2") ko pakadne wala pattern.
# Ye text search ke liye kachra hai, isliye hum ise hata dete hain.
DISCLAIMER_PATTERN = re.compile(
    r"DEMO\s*/\s*SYNTHETIC HOSPITAL DATA\s*[—–-]\s*NOT FOR REAL CLINICAL USE"
    r"(\s*\|\s*[A-Z]+-\d+\s*\|\s*Page\s*\d+\s*of\s*\d+)?",
    flags=re.IGNORECASE,
)


# ---------------------------------------------------------------
# 1) Metadata: filename se ID, title, department
# ---------------------------------------------------------------
def get_doc_info(path):
    """'EM-01_Emergency_Admission_Policy.pdf' -> ('EM-01', 'Emergency Admission Policy', 'Emergency')"""
    stem = Path(path).stem
    doc_id, _, rest = stem.partition("_")
    prefix = doc_id.split("-")[0]
    title = rest.replace("_", " ") or stem
    department = DEPARTMENTS.get(prefix, "General")
    return doc_id, title, department


# ---------------------------------------------------------------
# 2) Text cleaning
# ---------------------------------------------------------------
def clean_text(text):
    """Kachra hatao: disclaimer, footer, ====== lines, extra spaces, extra blank lines."""
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\u00a0", " ")
    text = DISCLAIMER_PATTERN.sub("", text)

    cleaned_lines = []
    for line in text.split("\n"):
        line = re.sub(r"[ \t]+", " ", line).strip()       # extra spaces hatao
        if re.fullmatch(r"[=\-_*]{3,}", line):             # ====== ya ------ wali lines hatao
            continue
        cleaned_lines.append(line)

    text = "\n".join(cleaned_lines)
    text = re.sub(r"\n{3,}", "\n\n", text)                 # 2 se zyada blank lines nahi
    return text.strip()


# ---------------------------------------------------------------
# 3) Har file type ka loader. Sab (page_number, text) ki list dete hain
# ---------------------------------------------------------------
def load_pdf(path):
    pages = []
    with fitz.open(str(path)) as pdf:
        for page_number, page in enumerate(pdf, start=1):
            # "blocks" = text ke tukde (lagbhag paragraphs). Har block ki line-breaks
            # ko space se badal dete hain taaki sentences beech se na tootein.
            blocks = page.get_text("blocks", sort=True)
            paragraphs = []
            for block in blocks:
                if block[6] != 0:          # 0 = text block, baaki images hain
                    continue
                paragraphs.append(block[4].replace("\n", " "))
            pages.append((page_number, clean_text("\n\n".join(paragraphs))))
    return pages


def load_docx(path):
    doc = Document(str(path))
    lines = []
    for para in doc.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style.name.startswith("List"):     # bullet points ko "- " se dikhao
            text = "- " + text
        lines.append(text)

    # Agar DOCX mein tables hon to unki rows bhi add kar do
    for table in doc.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            lines.append(" | ".join(cells))

    return [(1, clean_text("\n\n".join(lines)))]   # DOCX mein asli page nahi hota


def load_txt(path):
    text = Path(path).read_text(encoding="utf-8")
    return [(1, clean_text(text))]


def load_csv(path):
    # comment="#"          -> '#' se shuru hone wali metadata lines skip
    # dtype=str            -> sab kuch text ki tarah padho
    # keep_default_na=False -> "None" jaisa text NaN (khaali) na ban jaaye
    df = pd.read_csv(path, comment="#", dtype=str, keep_default_na=False)

    rows = []
    for _, row in df.iterrows():
        parts = []
        for column, value in row.items():
            value = str(value).strip()
            if value in ("", "-"):
                continue
            parts.append(f"{column.replace('_', ' ')}: {value}")
        rows.append(" | ".join(parts))

    return [(1, clean_text("\n\n".join(rows)))]    # har row ek alag paragraph


# ---------------------------------------------------------------
# 4) Ek file load karo, aur metadata jodo
# ---------------------------------------------------------------
def load_document(path):
    path = Path(path)
    doc_id, title, department = get_doc_info(path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        pages = load_pdf(path)
    elif extension == ".docx":
        pages = load_docx(path)
    elif extension == ".txt":
        pages = load_txt(path)
    elif extension == ".csv":
        pages = load_csv(path)
    else:
        raise ValueError(f"Unsupported file type: {extension}")

    records = []
    for page_number, text in pages:
        if not text:                                # khaali page skip
            continue
        records.append({
            "filename": path.name,
            "doc_id": doc_id,
            "title": title,
            "department": department,
            "file_type": extension.lstrip("."),
            "page": page_number,
            "text": text,
        })
    return records


# ---------------------------------------------------------------
# 5) Poore folder ki saari files load karo
# ---------------------------------------------------------------
def load_all_documents(docs_dir):
    all_records = []
    for path in sorted(Path(docs_dir).iterdir()):
        if not path.is_file() or path.suffix.lower() not in SUPPORTED_EXTENSIONS:
            continue
        try:
            all_records.extend(load_document(path))
        except Exception as error:                  # ek file kharab ho to baaki chalti rahein
            print(f"⚠️  {path.name} load nahi ho payi: {error}")
    return all_records


if __name__ == "__main__":
    import sys
    folder = sys.argv[1] if len(sys.argv) > 1 else "documents"
    records = load_all_documents(folder)
    print(f"Loaded {len(records)} page-records from {folder}")
